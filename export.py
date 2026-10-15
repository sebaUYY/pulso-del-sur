"""Pulso del Sur — archivos descargables y feeds para integrar con sistemas de gestión.
Escribe en site/datos/. Todo se regenera en cada actualización."""
import json, csv, os, calendar, datetime as dt
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

d = json.load(open("data.json"))
P = d["paises"]
ORDER = [k for k in ["UY", "AR", "BR", "CL", "PE", "CO", "MX"] if k in P]
OUT = "site/datos"
os.makedirs(OUT, exist_ok=True)
for sub in ["erp/netsuite", "erp/sap", "erp/odoo", "erp/business-central", "api"]:
    os.makedirs(f"{OUT}/{sub}", exist_ok=True)

CUR = {P[k]["moneda"]: k for k in ORDER}           # moneda -> país
BASES = ["USD"] + list(CUR)                          # moneda de la empresa (moneda funcional)
GEN = d["generado"]


# ---------------------------------------------------------------- helpers
def per_usd_on(cur, date):
    """Unidades de cur por 1 USD vigentes en date (último dato hábil <= date)."""
    if cur == "USD":
        return 1.0, None
    fx = P[CUR[cur]]["fx"]
    cand = [x for x in fx if x[0] <= date]
    return (cand[-1][1], cand[-1][0]) if cand else (None, None)


def latest_date():
    return max(P[k]["fx"][-1][0] for k in ORDER)


def cross(base, foreign, date):
    """Unidades de base por 1 unidad de foreign (vía USD si ninguna es USD)."""
    b, bdate = per_usd_on(base, date)
    f, fdate = per_usd_on(foreign, date)
    if b is None or f is None:
        return None, None
    # fecha efectiva = la del dato más antiguo de los dos (criterio conservador)
    return b / f, min(x for x in (bdate, fdate) if x)


def monthly_cross(base, foreign, m, key):
    """Cierre o promedio mensual cruzado vía USD."""
    def val(c):
        return 1.0 if c == "USD" else P[CUR[c]]["mensual"].get(m, {}).get(key)
    b, f = val(base), val(foreign)
    return (b / f) if (b and f) else None


def sap_q(r):
    """SAP guarda 5 decimales: para tasas chicas se usa factor (ej. 1000 ARS = x USD)."""
    fac = 1
    while r * fac < 1:
        fac *= 10
    return round(r * fac, 5), fac


def wcsv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


LD = latest_date()
months = sorted(set().union(*[set(P[k]["mensual"]) for k in ORDER]))
months = [m for m in months if m >= "2024-01"]
last_month = months[-1]


def month_end(m):
    y, mm = int(m[:4]), int(m[5:])
    return f"{m}-{calendar.monthrange(y, mm)[1]:02d}"

# ---------------------------------------------------------------- genéricos
# 1) último dato
latest = []
for k in ORDER:
    p = P[k]
    latest.append({"pais": p["nombre"], "moneda": p["moneda"], "fecha": p["fx"][-1][0],
                   "moneda_local_por_usd": p["fx"][-1][1], "usd_por_moneda_local": round(1 / p["fx"][-1][1], 10),
                   "fuente": p["fx_fuente"]})
wcsv(f"{OUT}/ultimo.csv", list(latest[0].keys()), [list(r.values()) for r in latest])
json.dump({"generado": GEN, "unidad": "moneda local por 1 USD", "datos": latest}, open(f"{OUT}/api/ultimo.json", "w"), ensure_ascii=False, indent=1)

# 2) histórico diario (formato largo)
diario = [(f, P[k]["nombre"], P[k]["moneda"], v) for k in ORDER for f, v in P[k]["fx"]]
wcsv(f"{OUT}/diario.csv", ["fecha", "pais", "moneda", "moneda_local_por_usd"], diario)
for k in ORDER:
    json.dump({"moneda": P[k]["moneda"], "pais": P[k]["nombre"], "fuente": P[k]["fx_fuente"], "unidad": "moneda local por 1 USD",
               "serie": [{"fecha": f, "valor": v} for f, v in P[k]["fx"]]},
              open(f"{OUT}/api/{P[k]['moneda'].lower()}.json", "w"), ensure_ascii=False)

# 3) cierre y promedio mensual
mens = []
for k in ORDER:
    for m, v in sorted(P[k]["mensual"].items()):
        if m >= "2024-01":
            mens.append((m, P[k]["nombre"], P[k]["moneda"], v["fecha_cierre"], v["cierre"], v["promedio"], v["dias"]))
wcsv(f"{OUT}/mensual.csv", ["mes", "pais", "moneda", "fecha_cierre", "cierre", "promedio_dias_habiles", "dias_habiles"], mens)
json.dump({"generado": GEN, "unidad": "moneda local por 1 USD",
           "datos": [dict(zip(["mes", "pais", "moneda", "fecha_cierre", "cierre", "promedio", "dias_habiles"], r)) for r in mens]},
          open(f"{OUT}/api/mensual.json", "w"), ensure_ascii=False)

# 4) Excel con todo (diario + mensual + último)
wb = Workbook()
hdr_font, hdr_fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F6B6F")
for title, header, rows in [("Último", ["País", "Moneda", "Fecha", "Moneda local por USD", "USD por moneda local", "Fuente"],
                             [list(r.values()) for r in latest]),
                            ("Mensual", ["Mes", "País", "Moneda", "Fecha cierre", "Cierre", "Promedio (días hábiles)", "Días hábiles"], mens),
                            ("Diario", ["Fecha", "País", "Moneda", "Moneda local por USD"], diario)]:
    ws = wb.active if title == "Último" else wb.create_sheet(title)
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([12, 12, 10, 14, 16, 22, 14][:len(header)]):
        ws.column_dimensions[chr(65 + i)].width = w
ws = wb.create_sheet("Notas")
for line in [f"Pulso del Sur · generado {GEN}", "Unidad: moneda local por 1 dólar estadounidense.",
             "Cierre = último día hábil con dato del mes. Promedio = media simple de días hábiles.",
             "Fuentes: " + " | ".join(f'{P[k]["nombre"]}: {P[k]["fx_fuente"]}' for k in ORDER),
             "Información de fuentes públicas oficiales. No es recomendación de inversión."]:
    ws.append([line])
wb.save(f"{OUT}/pulso-del-sur-datos.xlsx")

# ---------------------------------------------------------------- formatos por sistema
# Cada archivo existe por "moneda de la empresa" (base). Las tasas entre monedas latinoamericanas se cruzan vía USD.
for base in BASES:
    foreigns = [c for c in ["USD"] + list(CUR) if c != base]
    b = base.lower()

    # --- NetSuite: Currency Exchange Rates import (tasa = unidades de base por 1 unidad de moneda extranjera)
    rows = []
    for f in foreigns:
        r, fd_ = cross(base, f, LD)
        if r:
            rows.append((base, f, round(r, 8), fd_))
    wcsv(f"{OUT}/erp/netsuite/ultimo_base-{b}.csv", ["Base Currency", "Currency", "Exchange Rate", "Effective Date"], rows)
    rows = []
    for m in months:
        for f in foreigns:
            for key, typ in [("cierre", "Cierre"), ("promedio", "Promedio")]:
                r = monthly_cross(base, f, m, key)
                if r:
                    eff = month_end(m) if key == "cierre" else f"{m}-01"   # cierre: último día del mes; promedio: primer día
                    rows.append((base, f, round(r, 8), eff, typ))
    wcsv(f"{OUT}/erp/netsuite/mensual_base-{b}.csv", ["Base Currency", "Currency", "Exchange Rate", "Effective Date", "Tipo (Pulso del Sur)"], rows)

    # --- SAP (TCURR): FROM = moneda extranjera, TO = moneda de la empresa, tasa directa (1 FROM = x TO)
    rows = []
    for f in foreigns:
        r, fd_ = cross(base, f, LD)
        if r:
            q, fac = sap_q(r)
            rows.append(("M", f, base, fd_.replace("-", ""), q, fac, 1))
    wcsv(f"{OUT}/erp/sap/ultimo_to-{b}.csv", ["RATE_TYPE", "FROM_CURR", "TO_CURRNCY", "VALID_FROM", "EXCH_RATE", "FROM_FACTOR", "TO_FACTOR"], rows)
    rows = []
    for m in months:
        for f in foreigns:
            for key in ["cierre", "promedio"]:
                r = monthly_cross(base, f, m, key)
                if r:
                    q, fac = sap_q(r)
                    # RATE_TYPE vacío: cada empresa tiene sus propios tipos para cierre y promedio; se completa al importar
                    rows.append(("", "Cierre" if key == "cierre" else "Promedio", f, base,
                                 (month_end(m) if key == "cierre" else f"{m}-01").replace("-", ""), q, fac, 1))
    wcsv(f"{OUT}/erp/sap/mensual_to-{b}.csv", ["RATE_TYPE", "TIPO_PULSO", "FROM_CURR", "TO_CURRNCY", "VALID_FROM", "EXCH_RATE", "FROM_FACTOR", "TO_FACTOR"], rows)

    # --- Odoo (res.currency.rate): rate = unidades de moneda extranjera por 1 unidad de la moneda de la empresa
    rows = []
    for f in foreigns:
        r, fd_ = cross(base, f, LD)
        if r:
            rows.append((f, fd_, round(1 / r, 12), round(r, 8)))
    wcsv(f"{OUT}/erp/odoo/ultimo_empresa-{b}.csv", ["currency_id", "name", "rate", "inverse_company_rate"], rows)
    hist = []
    for f in foreigns:
        src = P[CUR[f]]["fx"] if f != "USD" else P[CUR[base]]["fx"]
        for fecha, _ in src:
            r, _fd = cross(base, f, fecha)
            if r:
                hist.append((f, fecha, round(1 / r, 12), round(r, 8)))
    wcsv(f"{OUT}/erp/odoo/historico_empresa-{b}.csv", ["currency_id", "name", "rate", "inverse_company_rate"], hist)

    # --- Business Central: feed XML para el "Currency Exchange Rate Service" (se lee solo, todos los días)
    # Exchange Rate Amount = 1 unidad de moneda extranjera; Relational = su valor en moneda local de la empresa (LCY)
    items = []
    for f in foreigns:
        r, fd_ = cross(base, f, LD)
        if r:
            items.append(f"  <rate><currency_code>{f}</currency_code><starting_date>{fd_}</starting_date>"
                         f"<exchange_rate_amount>1</exchange_rate_amount><relational_exch_rate_amount>{round(r, 6)}</relational_exch_rate_amount></rate>")
    with open(f"{OUT}/erp/business-central/lcy-{b}.xml", "w", encoding="utf-8") as fh:
        fh.write(f'<?xml version="1.0" encoding="UTF-8"?>\n<rates lcy="{base}" generated="{escape(GEN)}" source="Pulso del Sur">\n'
                 + "\n".join(items) + "\n</rates>\n")

# índice de archivos (para la página y para quien quiera automatizar)
files = sorted(os.path.relpath(os.path.join(r, f), "site") for r, _, fs in os.walk(OUT) for f in fs)
json.dump({"generado": GEN, "archivos": files}, open(f"{OUT}/indice.json", "w"), ensure_ascii=False, indent=1)
print(len(files), "archivos en", OUT)
